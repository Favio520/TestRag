from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import Any, Callable, List, Optional

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from openpyxl import load_workbook
from unstructured.partition.md import partition_md

logger = logging.getLogger(__name__)


def read_text_file(file_path: Path) -> str:
    try:
        return file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return file_path.read_text(encoding="latin-1")


def split_markdown_sections(raw_text: str) -> List[dict[str, Any]]:
    sections: List[dict[str, Any]] = []
    current_heading = "Contenido general"
    current_lines: List[str] = []

    for line in raw_text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            if current_lines:
                text = "\n".join(current_lines).strip()
                if text:
                    sections.append({"text": text, "metadata": {"section": current_heading}})
                current_lines = []
            current_heading = stripped.lstrip("#").strip() or "Contenido general"
            continue
        current_lines.append(line)

    if current_lines:
        text = "\n".join(current_lines).strip()
        if text:
            sections.append({"text": text, "metadata": {"section": current_heading}})

    return sections


def clean_tabular_row(row: dict[str, Any]) -> dict[str, str]:
    cleaned: dict[str, str] = {}
    for key, value in row.items():
        if key is None:
            continue
        normalized_key = str(key).strip()
        if not normalized_key or value is None:
            continue
        normalized_value = str(value).strip()
        if not normalized_value:
            continue
        cleaned[normalized_key] = normalized_value
    return cleaned


def resolve_record_id(row: dict[str, str]) -> Optional[str]:
    candidate_keys = (
        "id",
        "codigo",
        "code",
        "id_equipo",
        "id_sensor",
        "record_id",
        "nombre",
        "name",
    )
    lowered = {key.lower(): value for key, value in row.items()}
    for candidate in candidate_keys:
        value = lowered.get(candidate)
        if value:
            return value
    return None


def build_tabular_text(prefix: str, row: dict[str, str]) -> str:
    parts = [prefix]
    for key, value in row.items():
        parts.append(f"{key}: {value}.")
    return " ".join(parts).strip()


def extract_pdf_segments(file_path: Path) -> List[dict[str, Any]]:
    segments: List[dict[str, Any]] = []
    with fitz.open(file_path) as pdf_doc:
        for page_index, page in enumerate(pdf_doc, start=1):
            page_text = page.get_text("text")
            if page_text:
                segments.append(
                    {
                        "text": page_text.strip(),
                        "metadata": {"page_start": page_index, "page_end": page_index},
                    }
                )
    return segments


def extract_txt_segments(file_path: Path) -> List[dict[str, Any]]:
    text = read_text_file(file_path).strip()
    return [{"text": text, "metadata": {}}] if text else []


def extract_docx_segments(file_path: Path) -> List[dict[str, Any]]:
    doc = DocxDocument(str(file_path))
    sections: List[dict[str, Any]] = []
    current_heading = file_path.stem
    current_lines: List[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name.lower().startswith("heading"):
            if current_lines:
                body = "\n".join(current_lines).strip()
                if body:
                    sections.append({"text": body, "metadata": {"section": current_heading}})
                current_lines = []
            current_heading = text
            continue

        current_lines.append(text)

    if current_lines:
        body = "\n".join(current_lines).strip()
        if body:
            sections.append({"text": body, "metadata": {"section": current_heading}})

    return sections


def extract_md_segments(file_path: Path) -> List[dict[str, Any]]:
    raw_text = read_text_file(file_path)
    sections = split_markdown_sections(raw_text)
    if sections:
        return sections

    elements = partition_md(filename=str(file_path))
    lines = [el.text for el in elements if hasattr(el, "text") and el.text and el.text.strip()]
    text = "\n".join(lines).strip()
    return [{"text": text, "metadata": {}}] if text else []


def extract_csv_segments(file_path: Path) -> List[dict[str, Any]]:
    segments: List[dict[str, Any]] = []
    with file_path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        for row_number, row in enumerate(reader, start=2):
            clean_row = clean_tabular_row(dict(row))
            if not clean_row:
                continue
            text = build_tabular_text("Registro de tabla CSV.", clean_row)
            segments.append(
                {
                    "text": text,
                    "metadata": {
                        "table_name": file_path.stem,
                        "row_number": row_number,
                        "record_id": resolve_record_id(clean_row),
                    },
                }
            )
    return segments


def extract_xlsx_segments(file_path: Path) -> List[dict[str, Any]]:
    segments: List[dict[str, Any]] = []
    workbook = load_workbook(filename=str(file_path), read_only=True, data_only=True)
    try:
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [
                str(value).strip() if value is not None and str(value).strip() else f"columna_{idx + 1}"
                for idx, value in enumerate(rows[0])
            ]
            for row_number, values in enumerate(rows[1:], start=2):
                clean_row = clean_tabular_row(dict(zip(headers, values)))
                if not clean_row:
                    continue
                text = build_tabular_text(f"Registro de hoja {sheet_name}.", clean_row)
                segments.append(
                    {
                        "text": text,
                        "metadata": {
                            "sheet": sheet_name,
                            "table_name": file_path.stem,
                            "row_number": row_number,
                            "record_id": resolve_record_id(clean_row),
                        },
                    }
                )
    finally:
        workbook.close()
    return segments


EXTRACTORS: dict[str, Callable[[Path], List[dict[str, Any]]]] = {
    ".pdf": extract_pdf_segments,
    ".txt": extract_txt_segments,
    ".docx": extract_docx_segments,
    ".md": extract_md_segments,
    ".csv": extract_csv_segments,
    ".xlsx": extract_xlsx_segments,
}


def safe_extract_segments(file_path: Path) -> Optional[List[dict[str, Any]]]:
    extractor = EXTRACTORS.get(file_path.suffix.lower())
    if extractor is None:
        logger.warning("Formato no soportado, se omite: %s", file_path)
        return None

    try:
        segments = extractor(file_path)
    except Exception as exc:  # noqa: BLE001
        logger.warning("No se pudo procesar %s: %s", file_path, exc)
        return None

    valid_segments = [
        segment for segment in segments if str(segment.get("text", "")).strip()
    ]
    if not valid_segments:
        logger.warning("Archivo vacio o sin texto util, se omite: %s", file_path)
        return None
    return valid_segments

